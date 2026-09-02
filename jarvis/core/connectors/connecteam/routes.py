"""Connecteam connector API routes.

Manual Excel import for Bilet de Invoire form submissions.
No API integration — Connecteam plan does not support webhooks/API.
"""

import hashlib
import logging
from datetime import datetime, date, time

from flask import request, jsonify, g
from flask_login import current_user

from . import connecteam_bp
from .services import ConnecteamService
from .config import BILET_INVOIRE_FORM_ID, BILET_INVOIRE_FORM_NAME
from core.utils.api_helpers import api_login_required, admin_required, safe_error_response
from core.roles.decorators import v2_permission_required

logger = logging.getLogger('jarvis.connecteam.routes')
service = ConnecteamService()


# ── Status ──

@connecteam_bp.route('/api/status', methods=['GET'])
@api_login_required
def get_status():
    """Get connector status summary."""
    try:
        status = service.get_status()
        return jsonify({'success': True, 'data': status})
    except Exception as e:
        return safe_error_response(e)


# ── Excel Import (admin) ──

@connecteam_bp.route('/api/import-excel', methods=['POST'])
@api_login_required
def import_excel():
    """Import Connecteam form submissions from Excel export.

    Expects multipart/form-data with 'file' field containing .xlsx file.
    """
    if not getattr(current_user, 'can_access_hr', False) and not getattr(current_user, 'can_access_settings', False):
        return jsonify({'success': False, 'error': 'HR access required'}), 403
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': 'No file uploaded'}), 400

    file = request.files['file']
    if not file.filename or not file.filename.endswith('.xlsx'):
        return jsonify({'success': False, 'error': 'File must be .xlsx format'}), 400

    try:
        import openpyxl
        from io import BytesIO

        wb = openpyxl.load_workbook(BytesIO(file.read()), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()

        if not rows:
            return jsonify({'success': True, 'data': {
                'rows_processed': 0, 'inserted': 0, 'skipped': 0,
                'users_created': 0, 'unmapped_names': [],
            }})

        result = service.import_from_rows(rows)
        return jsonify({'success': True, 'data': result})

    except Exception as e:
        logger.exception('Excel import error: %s', e)
        return safe_error_response(e)


# ── User Mapping (admin) ──

@connecteam_bp.route('/api/users', methods=['GET'])
@admin_required
def get_users():
    """List Connecteam users with mapping status."""
    active_only = request.args.get('active_only', 'true').lower() == 'true'
    try:
        users = service.repo.get_all_users(active_only)
        return jsonify({'success': True, 'data': users})
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/users/auto-map', methods=['POST'])
@admin_required
def auto_map_users():
    """Auto-map Connecteam users to JARVIS users by name."""
    try:
        result = service.auto_map_users()
        return jsonify(result)
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/users/mapping', methods=['PUT'])
@admin_required
def update_mapping():
    """Manually map a Connecteam user to a JARVIS user."""
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'error': 'Request body required'}), 400

    connecteam_user_id = data.get('connecteam_user_id')
    jarvis_user_id = data.get('jarvis_user_id')

    if not connecteam_user_id or not jarvis_user_id:
        return jsonify({'success': False, 'error': 'connecteam_user_id and jarvis_user_id required'}), 400

    try:
        service.update_user_mapping(connecteam_user_id, jarvis_user_id)
        return jsonify({'success': True, 'message': 'Mapping updated'})
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/users/mapping', methods=['DELETE'])
@admin_required
def remove_mapping():
    """Remove mapping for a Connecteam user."""
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'error': 'Request body required'}), 400

    connecteam_user_id = data.get('connecteam_user_id')
    if not connecteam_user_id:
        return jsonify({'success': False, 'error': 'connecteam_user_id required'}), 400

    try:
        service.remove_user_mapping(connecteam_user_id)
        return jsonify({'success': True, 'message': 'Mapping removed'})
    except Exception as e:
        return safe_error_response(e)


# ── Approvers (for leave request form) ──

@connecteam_bp.route('/api/approvers', methods=['GET'])
@api_login_required
def get_approvers():
    """Get potential approvers for the current user's leave request.

    Walks UP the company structure from the user's assigned node(s)
    to collect all responsables above them. Falls back to L0 company
    responsables if the user has no structure node assignment.

    ?scope=all → returns all active users (for free-select mode).
    """
    try:
        scope = request.args.get('scope', '')
        if scope == 'all':
            from core.base_repository import BaseRepository
            rows = BaseRepository().query_all(
                "SELECT id, name FROM users WHERE is_active = TRUE ORDER BY name"
            )
            all_users = [{'id': r['id'], 'name': r['name']} for r in rows]
            return jsonify({'success': True, 'data': all_users})
        approvers = service.repo.get_approvers_for_user(current_user.id)
        return jsonify({'success': True, 'data': approvers})
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/leave-schedule', methods=['GET'])
@api_login_required
def get_leave_schedule_route():
    """Sincron work window + day cap + Forms-managed content for the leave form."""
    from .services.leave_schedule import get_leave_schedule
    from forms.services.form_service import FormService
    try:
        fs = FormService()
        data = get_leave_schedule(current_user.id, request.args.get('date'),
                                  request.args.get('company'))
        data.update(fs.get_leave_form_config())  # reasons, event_hours_reason, labels, placeholders, visible, terms_text
        # Two Time Bank pools shown separately on the form: personal (may go
        # negative) and event (capped, never negative — the "Ore Libere din
        # Eveniment" reason draws it and can't exceed it).
        split = fs.get_time_bank_split(current_user.id)
        data['time_bank_balance'] = split['total']       # back-compat / pooled
        data['time_bank_personal'] = split['personal']
        data['time_bank_event'] = split['event']
        # Corectie Ore monthly quota (mobile + web) — for the counter + toggle gate.
        # Everyone may file corrections; managers/HR/admins are exempt from the cap.
        data['corrections_limit'] = FormService.CORRECTION_MONTHLY_LIMIT
        data['corrections_used'] = fs.count_user_corrections_this_month(current_user.id)
        data['corrections_exempt'] = _is_correction_limit_exempt(current_user)
        # {id, name} of the direct manager the empty-approver default routes to, so
        # the form can auto-select it as a named chip on open.
        data['default_approver'] = fs.get_default_leave_approver(current_user.id)
        return jsonify({'success': True, 'data': data})
    except Exception as e:
        return safe_error_response(e)


# ── Submissions (user-scoped) ──

@connecteam_bp.route('/api/submissions/employee/<int:user_id>', methods=['GET'])
@api_login_required
@v2_permission_required('hr', 'leave_permissions', 'view')
def get_employee_submissions(user_id):
    """Get leave permission submissions for an employee (V2 permission scoped)."""
    from core.organization.hr_utils import get_managed_employee_ids

    scope = getattr(g, 'permission_scope', 'all')

    # Allow own data always; for others, check scope
    if user_id != current_user.id:
        if scope == 'own':
            return jsonify({'success': False, 'error': 'Permission denied'}), 403
        if scope == 'department':
            managed_ids = get_managed_employee_ids(current_user.id)
            if user_id not in (managed_ids or []):
                return jsonify({'success': False, 'error': 'Permission denied'}), 403

    year = request.args.get('year', type=int)
    month = request.args.get('month', type=int)

    try:
        data = service.get_user_submissions(user_id, year, month)
        return jsonify({'success': True, 'data': data})
    except Exception as e:
        return safe_error_response(e)


def _is_correction_limit_exempt(user) -> bool:
    """Corectie Ore is open to everyone (max 2/month) — but managers, HR and admins
    are NOT subject to that monthly cap."""
    if getattr(user, 'is_admin', False) or getattr(user, 'can_access_settings', False) \
            or getattr(user, 'is_hr_manager', False):
        return True
    try:
        from core.organization.manager_utils import is_manager
        return bool(is_manager(getattr(user, 'id', None)))
    except Exception:
        return False


@connecteam_bp.route('/api/submissions/leave-permit', methods=['POST'])
@api_login_required
def create_leave_permit():
    """Create a Bilet de Invoire from the code-defined Invoire module form.

    Fields are owned by the frontend (no DB form schema); the submission is
    stored as a form_submission and routed through the approval engine
    (primary manager + optional second approver, either approves).
    """
    from forms.services.form_service import FormService, UserContext
    data = request.get_json(silent=True) or {}
    answers = data.get('answers') if isinstance(data.get('answers'), dict) else data
    user = UserContext(user_id=current_user.id, company=getattr(current_user, 'company', None))
    try:
        result = FormService().submit_leave_permit(
            answers, user, ip_address=request.remote_addr,
            correction_limit_exempt=_is_correction_limit_exempt(current_user))
    except Exception as e:
        return safe_error_response(e)
    if not result.success:
        return jsonify({'success': False, 'error': result.error}), result.status_code
    return jsonify({'success': True, 'data': result.data}), result.status_code


@connecteam_bp.route('/api/submissions/leave-permit/<int:submission_id>/cancel', methods=['POST'])
@api_login_required
def cancel_leave_permit_route(submission_id):
    """Cancel (or withdraw) the current user's own leave-permit submission.
    A motive is required (shown to the manager on a cancellation request)."""
    from core.connectors.connecteam.services import leave_permit_actions as lpa
    reason = ((request.get_json(silent=True) or {}).get('reason') or '').strip()
    if not reason:
        return jsonify({'success': False, 'error': 'Motivul este obligatoriu.'}), 400
    try:
        data = lpa.cancel_leave_permit(submission_id, current_user.id, reason=reason)
        return jsonify({'success': True, 'data': data})
    except PermissionError:
        return jsonify({'success': False, 'error': 'Not your leave request'}), 403
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 409
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/submissions/leave-permit/<int:submission_id>', methods=['GET'])
@api_login_required
def get_leave_permit_route(submission_id):
    """Full stored answers for the current user's own leave-permit (edit prefill)."""
    from core.connectors.connecteam.services import leave_permit_actions as lpa
    try:
        return jsonify({'success': True, 'data': lpa.get_leave_permit(submission_id, current_user.id)})
    except PermissionError:
        return jsonify({'success': False, 'error': 'Not your leave request'}), 403
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 404
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/submissions/leave-permit/<int:submission_id>', methods=['PATCH'])
@api_login_required
def update_leave_permit_route(submission_id):
    """Modify the current user's own pending leave-permit submission."""
    from core.connectors.connecteam.services import leave_permit_actions as lpa
    body = request.get_json(silent=True) or {}
    answers = body.get('answers') if isinstance(body.get('answers'), dict) else body
    try:
        data = lpa.update_leave_permit(submission_id, current_user.id, answers)
        return jsonify({'success': True, 'data': data})
    except PermissionError:
        return jsonify({'success': False, 'error': 'Not your leave request'}), 403
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 409
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/leave-approvals/pending', methods=['GET'])
@api_login_required
def get_pending_leave_approvals():
    """Leave requests awaiting the current user's approval (empty if not an approver)."""
    try:
        data = service.get_pending_leave_approvals(current_user.id)
        return jsonify({'success': True, 'data': data})
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/leave-approvals/<int:request_id>/decide', methods=['POST'])
@api_login_required
def decide_leave_approval(request_id):
    """Approve or reject a leave request the current user is an approver for."""
    data = request.get_json(silent=True) or {}
    decision = data.get('decision')
    if decision not in ('approved', 'rejected'):
        return jsonify({'success': False, 'error': "decision must be 'approved' or 'rejected'"}), 400
    try:
        service.decide_leave_approval(request_id, decision, current_user.id, comment=data.get('comment'))
        return jsonify({'success': True})
    except PermissionError:
        return jsonify({'success': False, 'error': 'Not an approver for this request'}), 403
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/submissions/recent', methods=['GET'])
@admin_required
def get_recent_submissions():
    """Get submissions across all users (admin only), filtered by year/month.

    Merges Connecteam imported submissions and JARVIS internal form submissions.
    """
    limit = min(request.args.get('limit', 500, type=int), 1000)
    year = request.args.get('year', type=int)
    month = request.args.get('month', type=int)
    view = request.args.get('view', 'active')
    if view not in ('active', 'archived', 'trashed'):
        view = 'active'
    try:
        data = service.get_all_submissions(year=year, month=month, limit=limit, view=view)
        return jsonify({'success': True, 'data': data})
    except Exception as e:
        return safe_error_response(e)


# ── HR-scoped leave management (admin Leave-Permits tab) ──
# Edit details, or move a leave between lifecycle states — archive (filed) /
# delete (Coș, 7-day auto-purge) / restore (active) — either source.
# source ∈ {'jarvis', 'connecteam'}; validated in the service.

@connecteam_bp.route('/api/hr/leaves/<source>/<int:entity_id>', methods=['PATCH'])
@admin_required
def hr_edit_leave(source, entity_id):
    """HR override edit of a leave's details (date/start/end/reason). Status untouched."""
    from core.connectors.connecteam.services import leave_permit_actions as lpa
    fields = request.get_json(silent=True) or {}
    try:
        data = lpa.hr_update_leave(source, entity_id, fields)
        return jsonify({'success': True, 'data': data})
    except LookupError:
        return jsonify({'success': False, 'error': 'Bilet inexistent'}), 404
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return safe_error_response(e)


def _hr_set_state(source, entity_id, state):
    """Shared handler for the archive/delete/restore lifecycle routes."""
    from core.connectors.connecteam.services import leave_permit_actions as lpa
    try:
        data = lpa.hr_set_lifecycle(source, entity_id, current_user.id, state)
        return jsonify({'success': True, 'data': data})
    except LookupError:
        return jsonify({'success': False, 'error': 'Bilet inexistent'}), 404
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/hr/leaves/<source>/<int:entity_id>/archive', methods=['POST'])
@admin_required
def hr_archive_leave(source, entity_id):
    """Archive (file) a leave — recoverable via restore, kept indefinitely."""
    return _hr_set_state(source, entity_id, 'archived')


@connecteam_bp.route('/api/hr/leaves/<source>/<int:entity_id>/delete', methods=['POST'])
@admin_required
def hr_delete_leave(source, entity_id):
    """Move a leave to Coș/Trash — recoverable for 7 days, then auto-purged."""
    return _hr_set_state(source, entity_id, 'trashed')


@connecteam_bp.route('/api/hr/leaves/<source>/<int:entity_id>/restore', methods=['POST'])
@admin_required
def hr_restore_leave(source, entity_id):
    """Restore a leave from Archive or Trash back to the active list."""
    return _hr_set_state(source, entity_id, 'active')


# ── CO Conversions ──

@connecteam_bp.route('/api/conversions', methods=['POST'])
@admin_required
def create_conversion():
    """Initiate a CO conversion request for an employee's leave permits."""
    from .repositories.conversion_repository import ConversionRepository
    from hr.co_balance.repository import CoBalanceRepository
    from core.approvals.engine import ApprovalEngine

    data = request.get_json(force=True)
    employee_user_id = data.get('employee_user_id')
    year = data.get('year')
    month = data.get('month')
    co_days = data.get('co_days_requested')
    approver_user_id = data.get('approver_user_id')
    submission_ids = data.get('submission_ids', [])

    if not all([employee_user_id, year, month, co_days, approver_user_id]):
        return jsonify({'success': False, 'error': 'Missing required fields'}), 400

    conv_repo = ConversionRepository()
    co_repo = CoBalanceRepository()

    # Check no pending conversion for this employee/month
    existing = conv_repo.get_for_employee_month(employee_user_id, year, month)
    if existing:
        return jsonify({'success': False, 'error': 'A pending conversion already exists for this employee/month'}), 409

    # Get employee's accumulated hours for this month
    subs = service.repo.get_recent_submissions(500, year=year, month=month)
    user_subs = [s for s in subs if s.get('mapped_jarvis_user_id') == employee_user_id]
    total_hours = sum(float(s.get('leave_hours') or 0) for s in user_subs)

    if total_hours <= 0:
        return jsonify({'success': False, 'error': 'No accumulated hours for this employee/month'}), 400

    # Default to all submissions for this employee/month if none specified
    if not submission_ids:
        submission_ids = [s['submission_id'] for s in user_subs]

    # Check CO balance
    balance = co_repo.get_for_user(employee_user_id, year)
    if not balance:
        return jsonify({'success': False, 'error': 'No CO balance found for this employee'}), 400

    # Get employee name for context
    from core.base_repository import BaseRepository
    user_row = BaseRepository().query_one("SELECT name FROM users WHERE id = %s", (employee_user_id,))
    employee_name = user_row['name'] if user_row else f'User #{employee_user_id}'

    month_label = datetime(year, month, 1).strftime('%B %Y')

    try:
        # Create conversion record (approval_request_id will be set after)
        conversion_id = conv_repo.create(
            employee_user_id=employee_user_id,
            year=year, month=month,
            total_accumulated_hours=total_hours,
            co_days_requested=co_days,
            approver_user_id=approver_user_id,
            requested_by=current_user.id,
            approval_request_id=None,
            submission_ids=submission_ids,
        )

        # Submit to approval engine
        engine = ApprovalEngine()
        result = engine.submit(
            entity_type='leave_permit_conversion',
            entity_id=conversion_id,
            context={
                'approver_user_id': approver_user_id,
                'title': f'CO Conversion: {employee_name} — {co_days} days ({month_label})',
                'employee_name': employee_name,
                'employee_user_id': employee_user_id,
                'co_days_requested': co_days,
                'total_hours': float(total_hours),
                'year': year,
                'month': month,
            },
            requested_by=current_user.id,
        )

        # Link approval request back to conversion
        if result and result.get('request_id'):
            conv_repo.set_approval_request_id(conversion_id, result['request_id'])

        conversion = conv_repo.get_by_id(conversion_id)
        return jsonify({'success': True, 'data': dict(conversion) if conversion else {'id': conversion_id}})

    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/conversions', methods=['GET'])
@api_login_required
def list_conversions():
    """List conversion requests filtered by year/month."""
    from .repositories.conversion_repository import ConversionRepository

    year = request.args.get('year', type=int)
    month = request.args.get('month', type=int)
    if not year or not month:
        return jsonify({'success': False, 'error': 'year and month required'}), 400

    try:
        data = ConversionRepository().get_for_month(year, month)
        return jsonify({'success': True, 'data': data})
    except Exception as e:
        return safe_error_response(e)


@connecteam_bp.route('/api/conversions/<int:conversion_id>', methods=['GET'])
@api_login_required
def get_conversion(conversion_id):
    """Get a single conversion detail."""
    from .repositories.conversion_repository import ConversionRepository

    try:
        data = ConversionRepository().get_by_id(conversion_id)
        if not data:
            return jsonify({'success': False, 'error': 'Not found'}), 404
        return jsonify({'success': True, 'data': dict(data)})
    except Exception as e:
        return safe_error_response(e)
