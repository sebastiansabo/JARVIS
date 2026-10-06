"""Time Bank Excel export — builds a two-sheet .xlsx (Balanțe + Tranzacții).

Pure data→bytes helper: the route fetches the (scoped/filtered) rows and hands
them here, so this stays DB-free and unit-testable.
"""
import io
from datetime import datetime, timedelta, timezone

import openpyxl
from openpyxl.styles import Font

try:  # py3.9+ stdlib; backport on older runtimes (mirrors core/checkin/service.py)
    from zoneinfo import ZoneInfo
except ImportError:  # pragma: no cover
    from backports.zoneinfo import ZoneInfo

_RO_TZ = ZoneInfo('Europe/Bucharest')  # DST-correct, matches the web `toLocaleString('ro-RO')`

_BALANCE_HEADERS = ['Nume', 'Companie', 'Departament', 'Balanță (h)',
                    'Sold Personal', 'Sold Eveniment', 'Pending']
_BALANCE_WIDTHS = [26, 22, 22, 12, 14, 15, 9]

_TX_HEADERS = ['Dată', 'Angajat', 'Companie', 'Tip', 'Ore',
               'Status', 'Descriere', 'Creat de', 'Aprobat de']
_TX_WIDTHS = [17, 26, 22, 22, 8, 12, 42, 22, 22]


def _fmt_dt(value):
    """Render a stored UTC timestamp as Europe/Bucharest 'YYYY-MM-DD HH:MM'.

    Accepts a datetime or an ISO string; naive datetimes are treated as UTC.
    Anything unparseable is returned as-is so the export never crashes on a row.
    """
    if value is None:
        return ''
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return value
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.astimezone(_RO_TZ).strftime('%Y-%m-%d %H:%M')
    return str(value)


def _num(value):
    try:
        return float(value) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def _write_sheet(ws, headers, widths, rows):
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = 'A2'
    for row in rows:
        ws.append(row)


def build_time_bank_workbook(balances, transactions):
    """Return .xlsx bytes with a Balanțe sheet and a Tranzacții sheet."""
    wb = openpyxl.Workbook()

    ws_bal = wb.active
    ws_bal.title = 'Balanțe'
    _write_sheet(ws_bal, _BALANCE_HEADERS, _BALANCE_WIDTHS, [
        [
            b.get('name') or '',
            b.get('company') or '',
            b.get('department') or '',
            _num(b.get('balance')),
            _num(b.get('personal_balance')),
            _num(b.get('event_balance')),
            int(b.get('pending_count') or 0),
        ]
        for b in balances
    ])

    ws_tx = wb.create_sheet('Tranzacții')
    _write_sheet(ws_tx, _TX_HEADERS, _TX_WIDTHS, [
        [
            _fmt_dt(t.get('created_at')),
            t.get('employee_name') or '',
            t.get('employee_company') or '',
            t.get('tx_type') or '',
            _num(t.get('amount')),
            t.get('status') or '',
            t.get('description') or '',
            t.get('created_by_name') or '',
            t.get('approved_by_name') or '',
        ]
        for t in transactions
    ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
