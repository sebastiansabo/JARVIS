"""Monthly Foaie de Parcurs — Combustibil/Energie section is a reimbursement
(decontare) table:

  Km efectuați · Total alimentat · Normă consum/100 km · Normă proprie/km ·
  Total litri|kWh decontabili · Cost per litru|kWh · Valoare decontabilă

Decontabili = everything alimentat (Σ litri/kWh); Valoare decontabilă = Σ lei
(value of the alimentări at the price paid). Shown on BOTH HTML/PDF and Excel,
for the Combustibil (l) and Energie (kWh) sections.
"""
import os

os.environ.setdefault('DATABASE_URL', 'postgresql://localhost/defaultdb')

import foi_parcurs.services.route_sheet_service as rss

# 41.43 + 16.89 = 58.32 l alimentat / 392.34 + 161.64 = 553.98 lei over 843 km.
ENTRIES_L = [
    {'date': '2026-08-03', 'bon': '001', 'liters': 41.43, 'lei': 392.34},
    {'date': '2026-08-13', 'bon': '001', 'liters': 16.89, 'lei': 161.64},
]


# ── HTML (PDF) — the 7 decontare rows ────────────────────────────────────────
def test_fuel_html_has_all_seven_rows():
    html = rss._fuel_section_html('l', 6.5, ENTRIES_L, 843)
    assert 'Km efectuați</td><td>843 km' in html
    assert 'Total alimentat</td><td>58.32 l' in html
    assert 'Normă consum / 100 km</td><td>6.5 l/100 km' in html
    assert 'Normă proprie / km</td><td>0.065 l/km' in html
    assert 'Total litri decontabili</td><td>58.32 l' in html
    assert 'Cost per litru</td><td>9.5 lei/l' in html
    assert 'Valoare decontabilă</td><td>553.98 lei' in html


def test_fuel_html_no_deviation_or_old_rows():
    # The ±10% deviation alert and the old normat/efectiv/cost-total rows are gone.
    html = rss._fuel_section_html('l', 6.5, ENTRIES_L, 843)
    assert 'Deviație' not in html
    assert 'Consum normat' not in html
    assert 'Consum efectiv' not in html
    assert 'prag' not in html


def test_energy_html_uses_kwh_labels_and_units():
    html = rss._fuel_section_html('kWh', 17.0, [{'liters': 17.5, 'lei': 35}], 100)
    assert 'Total alimentat</td><td>17.5 kWh' in html
    assert 'Total kWh decontabili</td><td>17.5 kWh' in html
    assert 'Cost per kWh</td><td>2 lei/kWh' in html      # 35/17.5 = 2
    assert 'Normă proprie / km</td><td>0.17 kWh/km' in html
    assert 'Valoare decontabilă</td><td>35 lei' in html


def test_fuel_html_zero_when_no_alimentari():
    html = rss._fuel_section_html('l', 6.5, [], 843)
    assert 'Total alimentat</td><td>0 l' in html
    assert 'Total litri decontabili</td><td>0 l' in html
    assert 'Valoare decontabilă</td><td>0 lei' in html
    # No purchases → no average price.
    assert 'Cost per litru</td><td>—' in html


def test_fuel_html_dashes_when_no_norma():
    html = rss._fuel_section_html('l', None, ENTRIES_L, 843)
    assert 'Normă consum / 100 km</td><td>—' in html
    assert 'Normă proprie / km</td><td>—' in html
    # Alimentat/decontabili/valoare still computed from receipts.
    assert 'Total alimentat</td><td>58.32 l' in html
    assert 'Valoare decontabilă</td><td>553.98 lei' in html


# ── XLSX — same 7 rows, numeric cells ────────────────────────────────────────
def _xlsx_labels(norma, entries, km, unit='l'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active
    bold = Font(bold=True); fill = PatternFill('solid', fgColor='1A1A2E')
    rss._xlsx_fuel_section(ws, 1, unit, norma, entries, km, bold, fill, bold)
    out = {}
    for r in range(1, ws.max_row + 1):
        label = str(ws.cell(row=r, column=1).value or '')
        if label:
            out[label] = ws.cell(row=r, column=2).value
    return out


def test_xlsx_has_decontare_rows_with_values():
    labels = _xlsx_labels(6.5, ENTRIES_L, 843)
    # values are real numbers (not formatted strings) for spreadsheet math
    assert any(k.startswith('Km efectuați') and v == 843 for k, v in labels.items())
    assert any(k.startswith('Total alimentat') and v == 58.32 for k, v in labels.items())
    assert any(k.startswith('Normă proprie') and v == 0.065 for k, v in labels.items())
    assert any('decontabili' in k and v == 58.32 for k, v in labels.items())
    assert any(k.startswith('Valoare decontabilă') and v == 553.98 for k, v in labels.items())


def test_xlsx_has_no_deviation_row():
    labels = _xlsx_labels(6.5, ENTRIES_L, 843)
    assert not any(k.startswith('Deviație') for k in labels)


# ── Empty consumption sections are still hidden ──────────────────────────────
def _skeleton(fuel, fuel_type='Hybrid'):
    data = {
        'company': {'id': 1, 'name': 'Co', 'prestator': ''},
        'vehicle': {'vin': 'V', 'make': 'MG', 'model': 'HS', 'fuel_type': fuel_type,
                    'brand': 'MG Motor', 'category': 'M1', 'registration_number': 'R'},
        'period': {'year': 2026, 'month': 8, 'label': 'August 2026'},
        'trips': [], 'totals': {'km': 100, 'km_start': 0, 'km_end': 100, 'sessions': 0, 'clients': 0},
        'fuel': fuel, 'signatures': {},
    }
    return rss._skeleton_html(data, {'summary': '', 'trips': {}})


def test_plain_hybrid_shows_fuel_only():
    html = _skeleton({'norma': 7.5, 'norma_energie': 17.0,
                      'alimentari': [{'liters': 12, 'lei': 100, 'unit': 'l'}]}, fuel_type='Hybrid')
    assert 'fuel-title">Combustibil' in html
    assert 'fuel-title">Energie' not in html


def test_plugin_hybrid_shows_both_sections():
    html = _skeleton({'norma': 7.5, 'norma_energie': 17.0,
                      'alimentari': [{'liters': 12, 'lei': 100, 'unit': 'l'},
                                     {'liters': 8, 'lei': 20, 'unit': 'kWh'}]},
                     fuel_type='Plug-in Hybrid')
    assert 'fuel-title">Combustibil' in html
    assert 'fuel-title">Energie' in html


def test_empty_fuel_section_hidden():
    html = _skeleton({'norma': None, 'norma_energie': None, 'alimentari': []}, fuel_type='Benzina')
    assert 'fuel-title">Combustibil' not in html
