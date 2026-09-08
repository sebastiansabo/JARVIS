"""CSV / formula injection hardening for the deterministic Excel foaie.

User-controlled free text (Locul/Scopul override, Comentariu, event name, driver
names, alimentări bon/date) is written into .xlsx cells. A cell starting with
= + - @ is executed as a formula by Excel/Sheets, so such values must be
neutralized (leading apostrophe) — while numeric cells stay numeric.
"""
import io
import os

os.environ.setdefault('DATABASE_URL', 'postgresql://localhost/defaultdb')

import foi_parcurs.services.route_sheet_service as rss


def test_xl_safe_prefixes_formula_triggers():
    assert rss._xl_safe('=SUM(A1)') == "'=SUM(A1)"
    assert rss._xl_safe('+1') == "'+1"
    assert rss._xl_safe('-2') == "'-2"
    assert rss._xl_safe('@cmd') == "'@cmd"


def test_xl_safe_leaves_plain_text_and_numbers():
    assert rss._xl_safe('Service Brașov') == 'Service Brașov'
    assert rss._xl_safe('') == ''
    assert rss._xl_safe(42) == 42
    assert rss._xl_safe(3.5) == 3.5
    assert rss._xl_safe(None) is None


def test_render_xlsx_neutralizes_formula_in_scop_and_driver(monkeypatch):
    from openpyxl import load_workbook
    data = {
        'vehicle': {'vin': 'V', 'make': 'VW', 'model': 'Golf', 'fuel_type': 'Diesel',
                    'category': '', 'registration_number': 'B1'},
        'company': {'id': None, 'name': 'ACME', 'prestator': ''},
        'period': {'year': 2026, 'month': 8, 'label': 'august 2026'},
        'trips': [{'id': 1, 'plecare': '01.08.2026 10:00', 'sosire': '01.08.2026 12:00',
                   'km_start': 10, 'km_end': 20, 'distance_km': 10, 'is_td': False,
                   'driver': '=HYPERLINK("http://evil")', 'traseu': '=cmd|calc', 'itinerary': '',
                   'route_type': 'TD'}],
        'totals': {'km': 10, 'km_start': 10, 'km_end': 20, 'sessions': 1, 'clients': 1},
    }
    monkeypatch.setattr(rss, 'aggregate_month', lambda *a, **k: data)
    monkeypatch.setattr(rss, '_scop_overrides', lambda *a, **k: {})
    monkeypatch.setattr(rss, '_veh_repo', type('V', (), {'get_by_vin': lambda self, vin: {'fuel_type': 'Diesel'}})())
    monkeypatch.setattr(rss, '_store', type('S', (), {'query_one': lambda self, *a, **k: {}})())

    ws = load_workbook(io.BytesIO(rss.render_xlsx('V', 2026, 8))).active
    scop = [ws.cell(row=r, column=3).value for r in range(1, ws.max_row + 1)]
    driver = [ws.cell(row=r, column=4).value for r in range(1, ws.max_row + 1)]
    assert any(isinstance(v, str) and v.startswith("'=cmd") for v in scop)
    assert any(isinstance(v, str) and v.startswith("'=HYPERLINK") for v in driver)
