"""Time Bank Excel export: the workbook builder is a pure function (no DB) so we
can assert sheet layout, headers, values and Europe/Bucharest date rendering."""
import io
from datetime import datetime, timezone

import openpyxl

from hr.time_bank.export import build_time_bank_workbook


def _load(xlsx_bytes):
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def test_workbook_has_two_sheets_with_headers():
    wb = _load(build_time_bank_workbook(balances=[], transactions=[]))
    assert wb.sheetnames == ['Balanțe', 'Tranzacții']
    bal = [c.value for c in wb['Balanțe'][1]]
    tx = [c.value for c in wb['Tranzacții'][1]]
    assert bal == ['Nume', 'Companie', 'Departament', 'Balanță (h)',
                   'Sold Personal', 'Sold Eveniment', 'Pending']
    assert tx == ['Dată', 'Angajat', 'Companie', 'Tip', 'Ore',
                  'Status', 'Descriere', 'Creat de', 'Aprobat de']


def test_balance_row_values():
    balances = [{
        'name': 'Roman Paul-Gheorghe', 'company': 'Autoworld', 'department': 'Service',
        'balance': -15.0, 'personal_balance': -15.0, 'event_balance': 0.0, 'pending_count': 0,
    }]
    wb = _load(build_time_bank_workbook(balances=balances, transactions=[]))
    row = [c.value for c in wb['Balanțe'][2]]
    assert row == ['Roman Paul-Gheorghe', 'Autoworld', 'Service', -15.0, -15.0, 0.0, 0]


def test_transaction_date_rendered_in_romania_local_time():
    """06:12 UTC in October (DST) → 09:12 Europe/Bucharest, matching the web view."""
    transactions = [{
        'created_at': datetime(2026, 10, 5, 6, 12, tzinfo=timezone.utc),
        'employee_name': 'Roman Paul-Gheorghe', 'employee_company': 'Autoworld',
        'tx_type': 'manual_debit', 'amount': -4.0, 'status': 'approved',
        'description': 'BV Connecteam', 'created_by_name': 'Mates', 'approved_by_name': 'Mates',
    }]
    wb = _load(build_time_bank_workbook(balances=[], transactions=transactions))
    row = [c.value for c in wb['Tranzacții'][2]]
    assert row[0] == '2026-10-05 09:12'
    assert row[4] == -4.0
    assert row[3] == 'manual_debit'


def test_naive_and_none_dates_are_safe():
    transactions = [
        {'created_at': None, 'employee_name': 'X', 'amount': 1.0},
        {'created_at': '2026-01-15 10:00:00+00', 'employee_name': 'Y', 'amount': 2.0},
    ]
    wb = _load(build_time_bank_workbook(balances=[], transactions=transactions))
    rows = list(wb['Tranzacții'].iter_rows(min_row=2, values_only=True))
    assert not rows[0][0]              # None date → blank cell (openpyxl reads '' as None)
    assert rows[1][0] == '2026-01-15 12:00'  # Jan → UTC+2 (no DST)
