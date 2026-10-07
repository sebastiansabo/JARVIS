"""Tests for the project-leads webhook feature (Phase 1).

Covers:
- marketing.services.webhook_token  — token generation + hashing (pure)
- marketing.services.lead_intake     — webhook payload normalization (pure)
- marketing.repositories.webhook_repo.ProjectWebhookRepository (mocked DB)
- marketing.repositories.lead_repo.ProjectLeadRepository (mocked DB)
- marketing.routes.leads_webhook     — public intake endpoint (minimal Flask app)
"""
import os

os.environ.setdefault('DATABASE_URL', 'postgresql://test:test@localhost:5432/test')

import sys
from unittest.mock import patch, MagicMock

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'jarvis'))

_B = 'core.base_repository'


def _mock_db():
    return MagicMock(), MagicMock()


# ═══════════════════════════════════════════════
# webhook_token  (pure helpers)
# ═══════════════════════════════════════════════

class TestWebhookToken:

    def test_generate_token_shape(self):
        from marketing.services import webhook_token
        plaintext, token_hash, prefix = webhook_token.generate_token()
        assert plaintext.startswith('whk_')
        assert len(token_hash) == 64            # sha256 hex
        assert all(c in '0123456789abcdef' for c in token_hash)
        assert prefix and plaintext.startswith(prefix)

    def test_hash_token_matches_generated(self):
        from marketing.services import webhook_token
        plaintext, token_hash, _ = webhook_token.generate_token()
        assert webhook_token.hash_token(plaintext) == token_hash

    def test_hash_token_is_deterministic_and_distinct(self):
        from marketing.services import webhook_token
        assert webhook_token.hash_token('whk_abc') == webhook_token.hash_token('whk_abc')
        assert webhook_token.hash_token('whk_abc') != webhook_token.hash_token('whk_xyz')

    def test_hash_token_ignores_surrounding_whitespace(self):
        from marketing.services import webhook_token
        assert webhook_token.hash_token('  whk_abc  ') == webhook_token.hash_token('whk_abc')

    def test_generated_tokens_are_unique(self):
        from marketing.services import webhook_token
        a, _, _ = webhook_token.generate_token()
        b, _, _ = webhook_token.generate_token()
        assert a != b


# ═══════════════════════════════════════════════
# lead_intake.normalize_lead_payload  (pure)
# ═══════════════════════════════════════════════

class TestLeadIntake:

    def test_full_payload_normalized(self):
        from marketing.services.lead_intake import normalize_lead_payload
        data = {
            'contact_name': ' Ion Popescu ', 'phone': '+40722000000',
            'email': 'ion@example.com', 'company_name': 'ACME SRL', 'cui': 'RO123',
            'source': 'facebook_lead_ads', 'utm_source': 'facebook',
            'utm_campaign': 'summer', 'message': 'Interested', 'model_of_interest': 'XC60',
        }
        out = normalize_lead_payload(data)
        assert out['contact_name'] == 'Ion Popescu'          # trimmed
        assert out['email'] == 'ion@example.com'
        assert out['utm_source'] == 'facebook'
        assert out['raw_payload'] == data                    # original preserved

    def test_phone_only_is_valid(self):
        from marketing.services.lead_intake import normalize_lead_payload
        out = normalize_lead_payload({'phone': '0722000000'})
        assert out['phone'] == '0722000000'

    def test_email_only_is_valid(self):
        from marketing.services.lead_intake import normalize_lead_payload
        assert normalize_lead_payload({'email': 'a@b.ro'})['email'] == 'a@b.ro'

    def test_name_only_is_valid(self):
        from marketing.services.lead_intake import normalize_lead_payload
        assert normalize_lead_payload({'contact_name': 'Jane'})['contact_name'] == 'Jane'

    def test_empty_payload_rejected(self):
        from marketing.services.lead_intake import normalize_lead_payload
        with pytest.raises(ValueError):
            normalize_lead_payload({})

    def test_whitespace_only_identity_rejected(self):
        from marketing.services.lead_intake import normalize_lead_payload
        with pytest.raises(ValueError):
            normalize_lead_payload({'contact_name': '   ', 'phone': '', 'email': None})

    def test_non_dict_rejected(self):
        from marketing.services.lead_intake import normalize_lead_payload
        with pytest.raises(ValueError):
            normalize_lead_payload(['not', 'a', 'dict'])

    def test_unknown_keys_only_in_raw_payload(self):
        from marketing.services.lead_intake import normalize_lead_payload
        out = normalize_lead_payload({'phone': '0722', 'weird_key': 'keepme'})
        assert 'weird_key' not in out
        assert out['raw_payload']['weird_key'] == 'keepme'

    def test_long_value_truncated(self):
        from marketing.services.lead_intake import normalize_lead_payload
        out = normalize_lead_payload({'contact_name': 'x' * 5000})
        assert len(out['contact_name']) == 2000


# ═══════════════════════════════════════════════
# ProjectWebhookRepository  (mocked DB)
# ═══════════════════════════════════════════════

class TestProjectWebhookRepository:

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_create_returns_id(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchone.return_value = {'id': 7}

        from marketing.repositories.webhook_repo import ProjectWebhookRepository
        repo = ProjectWebhookRepository()
        result = repo.create(project_id=1, label='Zapier', token_hash='h', token_prefix='whk_abc', created_by=5)
        assert result == 7
        conn.commit.assert_called()

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_resolve_active_returns_row(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchone.return_value = {'id': 7, 'project_id': 1}

        from marketing.repositories.webhook_repo import ProjectWebhookRepository
        row = ProjectWebhookRepository().resolve_active('somehash')
        assert row == {'id': 7, 'project_id': 1}

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_resolve_active_none_when_missing(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchone.return_value = None

        from marketing.repositories.webhook_repo import ProjectWebhookRepository
        assert ProjectWebhookRepository().resolve_active('nope') is None

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_revoke_true_when_updated(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.rowcount = 1

        from marketing.repositories.webhook_repo import ProjectWebhookRepository
        assert ProjectWebhookRepository().revoke(project_id=1, webhook_id=7) is True

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_revoke_false_when_absent(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.rowcount = 0

        from marketing.repositories.webhook_repo import ProjectWebhookRepository
        assert ProjectWebhookRepository().revoke(project_id=1, webhook_id=999) is False


# ═══════════════════════════════════════════════
# ProjectLeadRepository  (mocked DB)
# ═══════════════════════════════════════════════

class TestProjectLeadRepository:

    def _normalized(self):
        from marketing.services.lead_intake import normalize_lead_payload
        return normalize_lead_payload({'contact_name': 'Jane', 'phone': '0722', 'x': 1})

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_create_returns_id(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchone.return_value = {'id': 42}

        from marketing.repositories.lead_repo import ProjectLeadRepository
        result = ProjectLeadRepository().create(
            project_id=1, payload=self._normalized(), received_via='webhook', webhook_id=7)
        assert result == 42
        conn.commit.assert_called()

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_update_status_true(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.rowcount = 1

        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().update(project_id=1, lead_id=42, status='contacted') is True

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_update_no_fields_false(self, mock_get_db, mock_get_cursor, mock_release):
        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().update(project_id=1, lead_id=42) is False

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_delete_true_when_removed(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.rowcount = 1

        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().delete(project_id=1, lead_id=42) is True

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_status_counts_maps_rows(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchall.return_value = [{'status': 'new', 'cnt': 3}, {'status': 'contacted', 'cnt': 1}]

        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().status_counts(project_id=1) == {'new': 3, 'contacted': 1}


# ═══════════════════════════════════════════════
# Phase 2 — ProjectLeadRepository additions
# ═══════════════════════════════════════════════

class TestProjectLeadRepositoryPhase2:

    def _norm(self):
        from marketing.services.lead_intake import normalize_lead_payload
        return normalize_lead_payload({'contact_name': 'Jane', 'phone': '0722'})

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_create_accepts_external_id_and_assignee(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchone.return_value = {'id': 99}

        from marketing.repositories.lead_repo import ProjectLeadRepository
        result = ProjectLeadRepository().create(
            project_id=1, payload=self._norm(), received_via='webhook',
            webhook_id=7, external_id='zap-abc', assigned_to=5)
        assert result == 99

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_update_assigns_user(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.rowcount = 1

        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().update(project_id=1, lead_id=42, assigned_to=5) is True

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_update_unassigns_user(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.rowcount = 1

        from marketing.repositories.lead_repo import ProjectLeadRepository
        # Explicit None must mean "unassign", not "leave unchanged".
        assert ProjectLeadRepository().update(project_id=1, lead_id=42, assigned_to=None) is True

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_update_truly_no_fields_false(self, mock_get_db, mock_get_cursor, mock_release):
        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().update(project_id=1, lead_id=42) is False

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_get_by_external_id(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchone.return_value = {'id': 12}

        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().get_by_external_id(1, 'zap-abc') == {'id': 12}

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_find_recent_duplicate_none(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.fetchone.return_value = None

        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().find_recent_duplicate(1, phone='0722', email=None, within_minutes=10) is None

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_find_recent_duplicate_needs_identity(self, mock_get_db, mock_get_cursor, mock_release):
        from marketing.repositories.lead_repo import ProjectLeadRepository
        # No phone and no email → nothing to match on → None, no query.
        assert ProjectLeadRepository().find_recent_duplicate(1, phone=None, email=None) is None

    @patch(f'{_B}.release_db')
    @patch(f'{_B}.get_cursor')
    @patch(f'{_B}.get_db')
    def test_mark_converted(self, mock_get_db, mock_get_cursor, mock_release):
        conn, cursor = _mock_db()
        mock_get_db.return_value = conn
        mock_get_cursor.return_value = cursor
        cursor.rowcount = 1

        from marketing.repositories.lead_repo import ProjectLeadRepository
        assert ProjectLeadRepository().mark_converted(project_id=1, lead_id=42, client_id=8, user_id=5) is True


# ═══════════════════════════════════════════════
# Phase 2 — convert lead → CRM client (resolve_or_create_client)
# ═══════════════════════════════════════════════

class _FakeClientRepo:
    # Mirrors the real ClientRepository finders' return shape: full rows (with
    # 'id') for find_by_phone / find_by_nr_reg. find_by_normalized_name is NOT
    # here on purpose — the service no longer reuses on a weak name-only match.
    def __init__(self, by_phone=None, by_nr=None, created_id=777):
        self._by_phone, self._by_nr = by_phone, by_nr
        self._created_id = created_id
        self.created_with = None

    def find_by_phone(self, phone):
        return self._by_phone

    def find_by_nr_reg(self, nr_reg):
        return self._by_nr

    def create_from_form(self, data):
        self.created_with = data
        return {'id': self._created_id}


class TestLeadConvert:

    def test_reuses_existing_by_phone(self):
        from marketing.services.lead_convert import resolve_or_create_client
        repo = _FakeClientRepo(by_phone={'id': 11})
        client_id, created = resolve_or_create_client(repo, {'contact_name': 'Jane', 'phone': '0722'})
        assert (client_id, created) == (11, False)
        assert repo.created_with is None   # did NOT create a duplicate

    def test_name_only_creates_new_no_weak_match(self):
        # A name with no phone/fiscal code is a WEAK identifier: never reuse,
        # always create a fresh client.
        from marketing.services.lead_convert import resolve_or_create_client
        repo = _FakeClientRepo(created_id=999)
        client_id, created = resolve_or_create_client(repo, {'contact_name': 'Jane Doe'})
        assert (client_id, created) == (999, True)
        assert repo.created_with['display_name'] == 'Jane Doe'

    def test_reuses_existing_by_fiscal_code(self):
        from marketing.services.lead_convert import resolve_or_create_client
        repo = _FakeClientRepo(by_nr={'id': 33})
        client_id, created = resolve_or_create_client(repo, {'company_name': 'ACME SRL', 'cui': 'RO123'})
        assert (client_id, created) == (33, False)

    def test_creates_person_when_no_match(self):
        from marketing.services.lead_convert import resolve_or_create_client
        repo = _FakeClientRepo(created_id=777)
        client_id, created = resolve_or_create_client(repo, {'contact_name': 'Jane', 'phone': '0722', 'email': 'j@x.ro'})
        assert (client_id, created) == (777, True)
        assert repo.created_with['client_type'] == 'person'
        assert repo.created_with['display_name'] == 'Jane'

    def test_creates_company_when_only_company_name(self):
        from marketing.services.lead_convert import resolve_or_create_client
        repo = _FakeClientRepo(created_id=888)
        client_id, created = resolve_or_create_client(repo, {'company_name': 'ACME SRL'})
        assert (client_id, created) == (888, True)
        assert repo.created_with['client_type'] == 'company'
        assert repo.created_with['display_name'] == 'ACME SRL'

    def test_raises_when_no_identity(self):
        from marketing.services.lead_convert import resolve_or_create_client
        with pytest.raises(ValueError):
            resolve_or_create_client(_FakeClientRepo(), {'message': 'no identity'})


# ═══════════════════════════════════════════════
# Phase 2 — leads Excel export (build_leads_workbook)
# ═══════════════════════════════════════════════

class TestLeadsExport:

    def _lead(self, **over):
        base = {
            'created_at': '2026-10-06T10:00:00', 'contact_name': 'Jane',
            'phone': '0722', 'email': 'j@x.ro', 'company_name': 'ACME',
            'cui': 'RO1', 'source': 'fb', 'utm_campaign': 'summer',
            'message': 'hi', 'model_of_interest': 'XC60', 'status': 'new',
            'assigned_to_name': 'Ana', 'status_notes': '',
        }
        base.update(over)
        return base

    def test_returns_xlsx_bytes(self):
        from marketing.services.leads_export import build_leads_workbook
        data = build_leads_workbook([self._lead()])
        assert isinstance(data, (bytes, bytearray)) and data[:2] == b'PK'  # xlsx = zip

    def test_header_and_row_present(self):
        import io
        import openpyxl
        from marketing.services.leads_export import build_leads_workbook
        wb = openpyxl.load_workbook(io.BytesIO(build_leads_workbook([self._lead()])))
        ws = wb.active
        assert ws.cell(row=1, column=2).value == 'Name'
        assert ws.cell(row=2, column=2).value == 'Jane'

    def test_formula_injection_neutralized(self):
        import io
        import openpyxl
        from marketing.services.leads_export import build_leads_workbook
        wb = openpyxl.load_workbook(io.BytesIO(build_leads_workbook([self._lead(message='=HACK()')])))
        ws = wb.active
        # Message is column 9; the leading '=' must be escaped to literal text.
        assert ws.cell(row=2, column=9).value == "'=HACK()"

    def test_empty_list_still_valid_workbook(self):
        import io
        import openpyxl
        from marketing.services.leads_export import build_leads_workbook
        wb = openpyxl.load_workbook(io.BytesIO(build_leads_workbook([])))
        assert wb.active.cell(row=1, column=1).value == 'Received'


# ═══════════════════════════════════════════════
# Public webhook intake route (minimal Flask app)
# ═══════════════════════════════════════════════

class TestLeadsWebhookRoute:

    URL = '/marketing/api/webhooks/leads'

    def _app(self):
        from flask import Flask
        from marketing.routes import leads_webhook as mod
        app = Flask(__name__)
        app.register_blueprint(mod.leads_webhook_bp, url_prefix='/marketing/api/webhooks')
        # Patch the module-level singletons so no real DB/rate state is touched.
        mod._webhook_repo = MagicMock()
        mod._lead_repo = MagicMock()
        # Default: no dedup hit, so a plain POST creates a new lead.
        mod._lead_repo.get_by_external_id.return_value = None
        mod._lead_repo.find_recent_duplicate.return_value = None
        mod._activity_repo = MagicMock()
        mod._rate_limiter = MagicMock()
        mod._rate_limiter.is_allowed.return_value = (True, 0)
        mod._ip_rate_limiter = MagicMock()
        mod._ip_rate_limiter.is_allowed.return_value = (True, 0)
        return app, mod

    def test_missing_token_401(self):
        app, mod = self._app()
        r = app.test_client().post(self.URL, json={'phone': '0722'})
        assert r.status_code == 401
        mod._lead_repo.create.assert_not_called()

    def test_invalid_token_401(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = None
        r = app.test_client().post(self.URL, json={'phone': '0722'},
                                   headers={'Authorization': 'Bearer whk_bad'})
        assert r.status_code == 401
        mod._lead_repo.create.assert_not_called()

    def test_valid_token_creates_lead_for_token_project(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        mod._lead_repo.create.return_value = 42
        r = app.test_client().post(
            self.URL,
            json={'contact_name': 'Jane', 'phone': '0722', 'source': 'fb', 'project_id': 999},
            headers={'Authorization': 'Bearer whk_good'})
        assert r.status_code == 201
        body = r.get_json()
        assert body['success'] is True and body['lead_id'] == 42
        # project comes from the TOKEN (3), never from the body (999)
        args, kwargs = mod._lead_repo.create.call_args
        assert (args and args[0] == 3) or kwargs.get('project_id') == 3
        mod._webhook_repo.touch_last_used.assert_called_once_with(7)

    def test_empty_lead_rejected_400(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        r = app.test_client().post(self.URL, json={'message': 'no identity here'},
                                   headers={'Authorization': 'Bearer whk_good'})
        assert r.status_code == 400
        mod._lead_repo.create.assert_not_called()

    def test_rate_limited_429(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        mod._rate_limiter.is_allowed.return_value = (False, 30)
        r = app.test_client().post(self.URL, json={'phone': '0722'},
                                   headers={'Authorization': 'Bearer whk_good'})
        assert r.status_code == 429
        mod._lead_repo.create.assert_not_called()

    def test_ip_throttle_blocks_before_token_lookup(self):
        """A per-IP flood is rejected 429 without ever hitting the token DB lookup."""
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        mod._lead_repo.create.return_value = 42
        mod._ip_rate_limiter.is_allowed.return_value = (False, 30)
        r = app.test_client().post(self.URL, json={'phone': '0722'},
                                   headers={'Authorization': 'Bearer whk_whatever'})
        assert r.status_code == 429
        mod._webhook_repo.resolve_active.assert_not_called()
        mod._lead_repo.create.assert_not_called()

    # ---- idempotency ----

    def test_external_id_duplicate_returns_existing_200(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        mod._lead_repo.get_by_external_id.return_value = {'id': 50}
        r = app.test_client().post(
            self.URL, json={'phone': '0722', 'external_id': 'zap-1'},
            headers={'Authorization': 'Bearer whk_good'})
        assert r.status_code == 200
        body = r.get_json()
        assert body['lead_id'] == 50 and body['duplicate'] is True
        mod._lead_repo.create.assert_not_called()

    def test_external_id_new_creates_201(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        mod._lead_repo.get_by_external_id.return_value = None
        mod._lead_repo.create.return_value = 51
        r = app.test_client().post(
            self.URL, json={'phone': '0722', 'external_id': 'zap-2'},
            headers={'Authorization': 'Bearer whk_good'})
        assert r.status_code == 201
        assert r.get_json()['lead_id'] == 51
        # external_id threaded into create
        _, kwargs = mod._lead_repo.create.call_args
        assert kwargs.get('external_id') == 'zap-2'

    def test_window_duplicate_returns_existing_200(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        mod._lead_repo.find_recent_duplicate.return_value = {'id': 52}
        r = app.test_client().post(
            self.URL, json={'phone': '0722', 'email': 'a@b.ro'},
            headers={'Authorization': 'Bearer whk_good'})
        assert r.status_code == 200
        body = r.get_json()
        assert body['lead_id'] == 52 and body['duplicate'] is True
        mod._lead_repo.create.assert_not_called()

    def test_idempotency_key_header_honored(self):
        app, mod = self._app()
        mod._webhook_repo.resolve_active.return_value = {'id': 7, 'project_id': 3}
        mod._lead_repo.get_by_external_id.return_value = {'id': 60}
        r = app.test_client().post(
            self.URL, json={'phone': '0722'},
            headers={'Authorization': 'Bearer whk_good', 'Idempotency-Key': 'hdr-1'})
        assert r.status_code == 200
        assert r.get_json()['lead_id'] == 60
        mod._lead_repo.get_by_external_id.assert_called_once()
        mod._lead_repo.create.assert_not_called()
